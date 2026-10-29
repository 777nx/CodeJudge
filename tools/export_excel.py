# -*- coding: utf-8 -*-
"""将 work/grades.json 导出为 Excel 成绩表。"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from paths import GRADES_JSON, GRADES_XLSX

GRADES = GRADES_JSON
OUT = GRADES_XLSX


def main() -> None:
    data = json.loads(GRADES.read_text(encoding="utf-8"))
    results = data.get("results", [])
    OUT.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩表"

    headers = [
        "序号",
        "座位号",
        "姓名",
        "来源文件",
        "题1-类名(20)",
        "题1-main注释(20)",
        "题1小计(40)",
        "题2-Student(10)",
        "题2-字段(10)",
        "题2-sayHello(10)",
        "题2-实例化(15)",
        "题2-运行(15)",
        "题2小计(60)",
        "总分(100)",
        "扣分/说明",
        "需复核",
        "复核原因",
        "置信度",
        "代码要点",
        "控制台输出",
        "截图数",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for i, r in enumerate(results, 1):
        s = r.get("scores") or {}
        a = int(s.get("A_classname", 0) or 0)
        b = int(s.get("B_main_comment", 0) or 0)
        c = int(s.get("C_student_class", 0) or 0)
        d = int(s.get("D_fields", 0) or 0)
        e = int(s.get("E_sayHello", 0) or 0)
        f = int(s.get("F_instantiate", 0) or 0)
        g = int(s.get("G_run", 0) or 0)
        q1 = a + b
        q2 = c + d + e + f + g
        total = int(r.get("total", q1 + q2) or 0)
        row = [
            i,
            r.get("seat") or "",
            r.get("name") or "",
            r.get("source") or "",
            a,
            b,
            q1,
            c,
            d,
            e,
            f,
            g,
            q2,
            total,
            r.get("deductions") or "",
            "是" if r.get("needs_review") else "否",
            r.get("review_reason") or "",
            r.get("confidence", ""),
            (r.get("code_summary") or "")[:500],
            (r.get("console_output") or "")[:300],
            r.get("image_count", 0),
        ]
        ws.append(row)
        if r.get("needs_review"):
            for col in range(1, len(headers) + 1):
                ws.cell(ws.max_row, col).fill = review_fill

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    widths = [6, 8, 10, 42, 12, 14, 10, 12, 10, 12, 12, 10, 10, 10, 28, 8, 24, 8, 36, 24, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    # 汇总页
    ws2 = wb.create_sheet("汇总")
    totals = [int(r.get("total", 0) or 0) for r in results]
    review_n = sum(1 for r in results if r.get("needs_review"))
    ws2.append(["考生人数", len(results)])
    ws2.append(["需复核人数", review_n])
    ws2.append(["平均分", round(sum(totals) / len(totals), 2) if totals else 0])
    ws2.append(["最高分", max(totals) if totals else 0])
    ws2.append(["最低分", min(totals) if totals else 0])
    ws2.append(["及格人数(>=60)", sum(1 for t in totals if t >= 60)])
    ws2.append(["优秀人数(>=90)", sum(1 for t in totals if t >= 90)])
    ws2.append(["零分人数", sum(1 for t in totals if t == 0)])

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
